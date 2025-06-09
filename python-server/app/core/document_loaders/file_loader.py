import os
import mimetypes
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Union

from .base import BaseLoader
from ..document import Document

def get_file_metadata(file_path: str) -> Dict[str, Any]:
    """Extract metadata from file."""
    file_path_obj = Path(file_path)
    return {
        "file_path": str(file_path_obj),
        "file_name": file_path_obj.name,
        "file_type": mimetypes.guess_type(file_path)[0],
        "file_size": os.path.getsize(file_path),
        "creation_date": datetime.fromtimestamp(
            file_path_obj.stat().st_ctime
        ).strftime("%Y-%m-%d"),
        "last_modified_date": datetime.fromtimestamp(
            file_path_obj.stat().st_mtime
        ).strftime("%Y-%m-%d"),
        "last_accessed_date": datetime.fromtimestamp(
            file_path_obj.stat().st_atime
        ).strftime("%Y-%m-%d"),
    }

class FileLoader(BaseLoader):
    """Load a single file."""
    
    def __init__(
        self, 
        file_path: str, 
        encoding: str = "utf-8", 
        errors: str = "ignore",
        filename_as_id: bool = False
    ):
        self.file_path = file_path
        self.encoding = encoding
        self.errors = errors
        self.filename_as_id = filename_as_id
        
    def load(self) -> List[Document]:
        """Load document from the file."""
        if not os.path.isfile(self.file_path):
            raise FileNotFoundError(f"File {self.file_path} does not exist")
            
        # Get file extension
        ext = os.path.splitext(self.file_path)[1].lower()
        
        # Extract metadata
        metadata = get_file_metadata(self.file_path)
        
        # Load content based on file type
        if ext == ".pdf":
            return self._load_pdf(metadata)
        elif ext == ".docx":
            return self._load_docx(metadata)
        elif ext == ".txt":
            return self._load_text(metadata)
        else:
            # Default handler for other file types
            return self._load_text(metadata)
            
    def _load_pdf(self, metadata: Dict[str, Any]) -> List[Document]:
        """Load PDF file."""
        try:
            import pypdf
        except ImportError:
            raise ImportError(
                "pypdf package not found, please install it with `pip install pypdf`"
            )
            
        documents = []
        try:
            with open(self.file_path, "rb") as file:
                pdf = pypdf.PdfReader(file)
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text.strip():  # Ignore empty pages
                        page_metadata = {
                            **metadata,
                            "page": i + 1,
                            "total_pages": len(pdf.pages)
                        }
                        doc = Document(page_content=text, metadata=page_metadata)
                        if self.filename_as_id:
                            doc.id_ = f"{self.file_path}:page{i+1}"
                        documents.append(doc)
                        
            return documents
        except Exception as e:
            raise ValueError(f"Error loading PDF file: {e}")
        
    def _load_docx(self, metadata: Dict[str, Any]) -> List[Document]:
        """Load DOCX file."""
        try:
            import docx
        except ImportError:
            raise ImportError(
                "python-docx package not found, please install it with `pip install python-docx`"
            )
            
        try:
            doc = docx.Document(self.file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            document = Document(page_content=text, metadata=metadata)
            if self.filename_as_id:
                document.id_ = self.file_path
            return [document]
        except Exception as e:
            raise ValueError(f"Error loading DOCX file: {e}")
            
    def _load_text(self, metadata: Dict[str, Any]) -> List[Document]:
        """Load text file."""
        try:
            with open(self.file_path, "r", encoding=self.encoding, errors=self.errors) as file:
                text = file.read()
                
            if not text.strip():  # Empty file
                return []
                
            document = Document(page_content=text, metadata=metadata)
            if self.filename_as_id:
                document.id_ = self.file_path
                
            return [document]
        except Exception as e:
            raise ValueError(f"Error loading text file {self.file_path}: {e}")


class MultiFileLoader(BaseLoader):
    """Load multiple files."""
    
    def __init__(
        self, 
        file_paths: List[str], 
        encoding: str = "utf-8", 
        errors: str = "ignore",
        filename_as_id: bool = False
    ):
        self.file_paths = file_paths
        self.encoding = encoding
        self.errors = errors
        self.filename_as_id = filename_as_id
        
    def load(self) -> List[Document]:
        """Load documents from multiple files."""
        documents = []
        
        for file_path in self.file_paths:
            try:
                loader = FileLoader(
                    file_path=file_path,
                    encoding=self.encoding,
                    errors=self.errors,
                    filename_as_id=self.filename_as_id
                )
                file_docs = loader.load()
                documents.extend(file_docs)
            except Exception as e:
                print(f"Error loading file {file_path}: {e}")
                
        return documents
